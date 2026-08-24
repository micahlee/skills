#!/usr/bin/env python3
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
FIXTURES = Path(__file__).resolve().parent / "fixtures"
NORMALIZED = FIXTURES / "chasen-monday.json"
SAMUEL_NORMALIZED = FIXTURES / "samuel-monday.json"
NEXT_DAY = FIXTURES / "next-day-chasen.json"


class AgendaToolTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="agenda-tools-test-")
        self.work = Path(self.temp.name)
        self.docx = self.work / "Chasen-Agenda-2026-08-24.docx"
        subprocess.run(
            [sys.executable, str(SCRIPTS / "build_agendas.py"), str(NORMALIZED), str(self.docx)],
            check=True,
        )

    def tearDown(self):
        self.temp.cleanup()

    def run_json(self, *arguments: str, expected: int = 0):
        completed = subprocess.run(
            [sys.executable, *arguments],
            check=False,
            capture_output=True,
            text=True,
        )
        self.assertEqual(completed.returncode, expected, completed.stderr or completed.stdout)
        stream = completed.stdout if completed.stdout.strip() else completed.stderr
        return json.loads(stream)

    def test_validator_accepts_compact_agenda(self):
        result = self.run_json(str(SCRIPTS / "validate_agenda.py"), str(NORMALIZED), str(self.docx))
        self.assertTrue(result["ok"])

    def test_validator_accepts_samuel_fixture(self):
        samuel_docx = self.work / "Samuel-Agenda-2026-08-24.docx"
        subprocess.run(
            [sys.executable, str(SCRIPTS / "build_agendas.py"), str(SAMUEL_NORMALIZED), str(samuel_docx)],
            check=True,
        )
        result = self.run_json(
            str(SCRIPTS / "validate_agenda.py"), str(SAMUEL_NORMALIZED), str(samuel_docx)
        )
        self.assertTrue(result["ok"])

    def test_empty_schedule_keeps_bible_and_exact_no_prep_line(self):
        empty_json = self.work / "empty.json"
        empty_docx = self.work / "empty.docx"
        data = json.loads(NORMALIZED.read_text())
        data["courses"] = []
        data["next_day_prep"] = []
        data["parent_report"]["incomplete"] = []
        empty_json.write_text(json.dumps(data))
        subprocess.run(
            [sys.executable, str(SCRIPTS / "build_agendas.py"), str(empty_json), str(empty_docx)],
            check=True,
        )
        result = self.run_json(
            str(SCRIPTS / "validate_agenda.py"), str(empty_json), str(empty_docx)
        )
        self.assertTrue(result["ok"])
        visible = "\n".join(cell.text for table in Document(empty_docx).tables for row in table.rows for cell in row.cells)
        self.assertIn("Bible", visible)
        self.assertIn("No special ClassReach prep notes found", visible)

    def test_validator_rejects_detail_in_school_table(self):
        document = Document(self.docx)
        document.tables[1].rows[0].cells[1].paragraphs[0].add_run(" - forbidden detail")
        bad = self.work / "bad.docx"
        document.save(bad)
        result = self.run_json(
            str(SCRIPTS / "validate_agenda.py"), str(NORMALIZED), str(bad), expected=1
        )
        self.assertFalse(result["ok"])
        self.assertTrue(any("School row 1" in error for error in result["errors"]))

    def test_validator_rejects_internally_inconsistent_counts(self):
        inconsistent = self.work / "inconsistent.json"
        data = json.loads(NORMALIZED.read_text())
        data["courses"][0]["assignment_count"] = 2
        inconsistent.write_text(json.dumps(data))
        result = self.run_json(
            str(SCRIPTS / "validate_agenda.py"), str(inconsistent), str(self.docx), expected=1
        )
        self.assertFalse(result["ok"])
        self.assertTrue(any("does not match 1 assignment records" in error for error in result["errors"]))

    def test_next_day_classifier_preserves_terse_task_and_material_sentence(self):
        result = self.run_json(
            str(SCRIPTS / "classify_next_day_prep.py"),
            str(NEXT_DAY),
            "--student",
            "Chasen",
        )
        self.assertFalse(result["no_special_preparation_allowed"])
        self.assertEqual(
            [(item["course"], item["text"]) for item in result["next_day_prep"]],
            [
                ("Grammar 7", "Bible"),
                ("History Transition", "Students should write the answers in their spiral notebook."),
            ],
        )
        classes = {item["course"]: item["classification"] for item in result["items"]}
        self.assertEqual(classes["Latin 1-7"], "ordinary_in_class")
        self.assertEqual(classes["Science"], "optional")

    def test_next_day_classifier_forbids_false_empty_claim(self):
        result = self.run_json(
            str(SCRIPTS / "classify_next_day_prep.py"),
            str(NEXT_DAY),
            "--student",
            "Chasen",
            "--expect-empty",
            expected=1,
        )
        self.assertFalse(result["ok"])

    def test_next_day_classifier_fails_closed_on_unknown_item_type(self):
        unknown = self.work / "unknown-next-day.json"
        fixture = json.loads(NEXT_DAY.read_text())
        fixture["students"][0]["sections"][0]["agendaItems"] = [
            {
                "type": "MysteryItem",
                "title": "Optional unclassified work",
                "details": "Optional: bring your workbook.",
                "meetingDay": False,
            }
        ]
        unknown.write_text(json.dumps(fixture))
        result = self.run_json(
            str(SCRIPTS / "classify_next_day_prep.py"),
            str(unknown),
            "--student",
            "Chasen",
            expected=1,
        )
        self.assertFalse(result["ok"])
        self.assertIn("unclassified item", result["error"])

    def test_publisher_preflights_entire_batch_before_copying(self):
        bad = self.work / "bad-for-batch.docx"
        document = Document(self.docx)
        document.tables[1].rows[0].cells[1].paragraphs[0].add_run(" - forbidden detail")
        document.save(bad)
        destination = self.work / "batch-published"
        result = self.run_json(
            str(SCRIPTS / "publish_agendas.py"),
            "--agenda",
            str(NORMALIZED),
            str(self.docx),
            "--agenda",
            str(NORMALIZED),
            str(bad),
            "--destination",
            str(destination),
            "--renderer",
            str(Path(__file__).resolve().parent / "fake_renderer.py"),
            expected=1,
        )
        self.assertFalse(result["ok"])
        self.assertFalse(destination.exists())

    def test_publisher_uses_immutable_revisions_and_manifests(self):
        destination = self.work / "published"
        renders = self.work / "renders"
        command = [
            str(SCRIPTS / "publish_agendas.py"),
            "--agenda",
            str(NORMALIZED),
            str(self.docx),
            "--destination",
            str(destination),
            "--renderer",
            str(Path(__file__).resolve().parent / "fake_renderer.py"),
            "--render-dir",
            str(renders),
        ]
        first = self.run_json(*command)
        first_output = first["outputs"][0]
        self.assertEqual(Path(first_output["published_docx"]).name, self.docx.name)
        self.assertTrue(Path(first_output["manifest"]).is_file())
        manifest = json.loads(Path(first_output["manifest"]).read_text())
        self.assertEqual(
            set(manifest),
            {
                "student",
                "date",
                "docx",
                "docx_sha256",
                "normalized_json_sha256",
                "page_count",
                "page_sha256",
                "reused_identical_file",
            },
        )
        self.assertNotIn("Bible", Path(first_output["manifest"]).read_text())

        alternate_dir = self.work / "alternate"
        alternate_dir.mkdir()
        alternate = alternate_dir / self.docx.name
        document = Document(self.docx)
        document.core_properties.subject = "hash-changing but layout-neutral revision"
        document.save(alternate)
        command[command.index(str(self.docx))] = str(alternate)
        second = self.run_json(*command)
        second_output = second["outputs"][0]
        self.assertEqual(Path(second_output["published_docx"]).name, f"{self.docx.stem}-r2.docx")
        self.assertTrue(Path(second_output["manifest"]).is_file())


if __name__ == "__main__":
    unittest.main()
