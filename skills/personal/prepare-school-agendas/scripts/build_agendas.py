#!/usr/bin/env python3
"""Build one printable student agenda from normalized, verified JSON."""

from __future__ import annotations

import argparse
import json
import re
from datetime import date as date_type
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INK = "111111"
MUTED = "555555"
LIGHT = "E7E7E7"
PALE = "F5F5F5"
WHITE = "FFFFFF"
GRID = "555555"
FONT = "Arial"
PAGE_WIDTH_DXA = 10656  # 7.4 in after 0.55 in side margins
CHECK_DXA = 560
TEXT_DXA = PAGE_WIDTH_DXA - CHECK_DXA


def set_run(run, size=10.2, bold=False, italic=False, color=INK, strike=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.strike = strike
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=55, start=100, bottom=55, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_borders(cell, color=GRID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def setup_table(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    indent = tbl_pr.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:w"), "100")
    indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_width(cell, width)
            set_cell_margins(cell)
            set_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def clear_paragraph(p):
    for child in list(p._p):
        p._p.remove(child)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    return p


def add_hyperlink(paragraph, text, url, *, bold=False, strike=False, size=9.5):
    if not url:
        run = paragraph.add_run(text)
        set_run(run, size=size, bold=bold, strike=strike)
        return run
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), INK)
    rpr.append(color)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if strike:
        rpr.append(OxmlElement("w:strike"))
    run.append(rpr)
    text_node = OxmlElement("w:t")
    if text[:1].isspace() or text[-1:].isspace():
        text_node.set(qn("xml:space"), "preserve")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_section_bar(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT)
    ppr.append(shd)
    run = p.add_run(title.upper())
    set_run(run, size=10.0, bold=True)
    return p


def add_checkbox_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    for row_data in rows:
        cells = table.add_row().cells
        checked = bool(row_data.get("checked"))
        p0 = clear_paragraph(cells[0].paragraphs[0])
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        marker = row_data.get("marker", "☒" if checked else "☐")
        set_run(p0.add_run(marker), size=14, bold=False)
        p = clear_paragraph(cells[1].paragraphs[0])
        if "parts" in row_data:
            for part in row_data["parts"]:
                add_hyperlink(p, part["text"], part.get("url", ""), bold=part.get("bold", False), strike=part.get("strike", False), size=part.get("size", 9.0))
        else:
            add_hyperlink(p, row_data["text"], row_data.get("url", ""), bold=row_data.get("bold", False), strike=row_data.get("strike", False), size=row_data.get("size", 9.5))
        if row_data.get("note"):
            r = p.add_run(row_data["note"])
            set_run(r, size=8.7, italic=True, color=MUTED)
    setup_table(table, [CHECK_DXA, TEXT_DXA])
    return table


def add_inline_checkbox_table(doc, rows):
    """Use one cell so LibreOffice cannot drop isolated checkbox-cell text."""
    table = doc.add_table(rows=0, cols=1)
    for row_data in rows:
        cell = table.add_row().cells[0]
        p = clear_paragraph(cell.paragraphs[0])
        set_run(p.add_run("☐  "), size=14)
        add_hyperlink(p, row_data["text"], row_data.get("url", ""), size=row_data.get("size", 9.5))
    setup_table(table, [PAGE_WIDTH_DXA])
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10.2)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0


def add_title(doc, student, target):
    display = target.strftime("%A, %B %-d, %Y")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(f"{student} - {display}"), size=17, bold=True)


def morning_rows():
    return [{"text": x} for x in ("Make your bed", "Pick up trash around your room", "Get dressed")]


def fine_arts_rows(student, weekday):
    if student == "Samuel":
        labels = ["Percussion Lesson", "Percussion Practice"] if weekday == "Monday" else ["Percussion Practice"]
    elif student == "Chasen":
        labels = ["Trumpet Lesson", "Trumpet Practice", "Theater Class", "Theater Practice"] if weekday == "Monday" else ["Trumpet Practice", "Theater Practice"]
    else:
        raise ValueError(f"unsupported student: {student}")
    return [{"text": f"{label} - Minutes: ____________________"} for label in labels]


def detail_for_print(details, *, limit=500, max_newlines=6):
    details = (details or "").strip()
    if not details:
        return ""
    details = re.sub(r"https?://\S+", "(link in ClassReach).", details)
    if len(details) <= limit and details.count("\n") <= max_newlines:
        return details.replace("\n", " ")
    return "View details in ClassReach"


def school_rows(data):
    courses = data.get("courses", [])
    if not courses or not any(int(course.get("assignment_count", 0)) for course in courses):
        checked_at = data.get("checked_at", "verification time unavailable")
        rows = [
            {"text": "No assignments listed in ClassReach", "bold": True, "note": f"  Checked on {checked_at}"},
            {"text": "Bible"},
        ]
    else:
        rows = []
        printed_bible = False
        for course in courses:
            count = int(course.get("assignment_count", 0))
            if count <= 0:
                continue
            if course.get("name") == "Bible":
                printed_bible = True
            complete = int(course.get("completed_count", 0)) >= count
            noun = "assignment" if count == 1 else "assignments"
            rows.append({
                "text": f"{course['name']} - {count} {noun}",
                "url": course.get("url", ""),
                "checked": complete,
                "strike": complete,
            })
        if not printed_bible:
            rows.append({"text": "Bible"})
    return rows


def afternoon_rows(data, weekday):
    rows = [{"text": "Tidy your room and help tidy the house"}]
    if weekday in {"Monday", "Wednesday"}:
        prep = data.get("next_day_prep", [])
        rows.append({"text": "Prepare for school tomorrow", "bold": True})
        if prep:
            for item in prep:
                rows.append({"text": f"{item.get('course', 'ClassReach')}: {item['text']}", "url": item.get("url", ""), "size": 9.2})
        else:
            rows.append({"text": "No special ClassReach prep notes found", "size": 9.0, "marker": ""})
        rows.append({"text": "Make your lunch"})
        if weekday == "Wednesday":
            rows.append({"text": "Gather the towels for laundry" if data["student"] == "Samuel" else "Gather the family trash"})
    else:
        rows.append({"text": "Put your school stuff away"})
    return rows


def additional_task_rows():
    return [{"text": "____________________________________________________________"} for _ in range(3)]


def add_parent_report(doc, data, target):
    doc.add_section(WD_SECTION.NEW_PAGE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run(f"{data['student']} - Parent Report"), size=17, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_run(p.add_run(target.strftime("%A, %B %-d, %Y")), size=10.5, color=MUTED)

    report = data.get("parent_report") or {}
    sections = [
        ("Incomplete work", report.get("incomplete", [])),
        ("Messages, discussions, and handouts", report.get("messages", [])),
        ("Schedule or agenda changes", report.get("changes", [])),
    ]
    for heading, items in sections:
        add_section_bar(doc, heading)
        if not items:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            set_run(p.add_run("Nothing to report"), size=10.2, italic=True, color=MUTED)
            continue
        rows = []
        if heading == "Incomplete work":
            counts = {}
            urls = {}
            for item in items:
                course = item.get("course") or "ClassReach"
                counts[course] = counts.get(course, 0) + 1
                urls.setdefault(course, item.get("url", ""))
            for course, count in counts.items():
                noun = "assignment" if count == 1 else "assignments"
                rows.append({"text": f"{course} - {count} {noun}", "url": urls[course], "size": 9.5})
            add_inline_checkbox_table(doc, rows)
            continue
        for item in items:
            text = item.get("text") or item.get("title") or "ClassReach item"
            course = item.get("course")
            if course:
                text = f"{course}: {text}"
            details = (item.get("details") or "").strip()
            detail_limit = 300
            newline_limit = 2
            if details and item.get("details_mode") != "link" and len(details) <= detail_limit and details.count("\n") <= newline_limit:
                text += f" - {details}"
            elif details:
                text += " - View details in ClassReach"
            rows.append({"text": text, "url": item.get("url", ""), "size": 9.5})
        add_inline_checkbox_table(doc, rows)


def build(data: dict[str, Any], output: Path, template: Path | None = None):
    student = data.get("student")
    if student not in {"Samuel", "Chasen"}:
        raise ValueError("student must be Samuel or Chasen")
    target = date_type.fromisoformat(data["date"])
    weekday = target.strftime("%A")
    if weekday not in {"Monday", "Wednesday", "Friday"}:
        raise ValueError("target date must be Monday, Wednesday, or Friday")
    if weekday in {"Monday", "Friday"} and "parent_report" not in data:
        raise ValueError("Monday and Friday require parent_report")

    doc = Document(template) if template else Document()
    if template:
        body = doc._element.body
        final_sect_pr = body.sectPr
        for child in list(body):
            if child is not final_sect_pr:
                body.remove(child)
    configure_document(doc)
    add_title(doc, student, target)

    add_section_bar(doc, "Morning Checklist")
    add_checkbox_table(doc, morning_rows())

    add_section_bar(doc, "School")
    add_checkbox_table(doc, school_rows(data))

    add_section_bar(doc, "Fine Arts")
    add_checkbox_table(doc, fine_arts_rows(student, weekday))

    add_section_bar(doc, "Physical Education")
    add_checkbox_table(doc, [{"text": "Activity: ________________________________    Minutes: __________"}])

    add_section_bar(doc, "Afternoon Checklist")
    add_checkbox_table(doc, afternoon_rows(data, weekday))

    add_section_bar(doc, "Additional Tasks")
    add_checkbox_table(doc, additional_task_rows())

    if weekday in {"Monday", "Friday"}:
        add_parent_report(doc, data, target)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = f"{student} Agenda - {target.isoformat()}"
    doc.core_properties.subject = "Printable school agenda"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--template", type=Path)
    args = parser.parse_args()
    with args.input.open("r", encoding="utf-8") as handle:
        data = json.load(handle)
    build(data, args.output, args.template)


if __name__ == "__main__":
    main()
